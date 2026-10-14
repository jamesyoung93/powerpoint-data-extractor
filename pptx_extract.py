"""Extract labeled datasets from PPTX/PPTM packages without Office or an AI service.

Python >=3.10. Dependencies: openpyxl>=3.1,<4, defusedxml>=0.7.1,<1.
Use extract_pptx(path_or_bytes), write_outputs(report, directory), and
to_spark_dataframe(spark, report). The companion notebook is self-contained.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import math
import posixpath
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import unquote

from defusedxml import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, coordinate_to_tuple

__version__ = "1.1.0"
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}
RID = "{" + NS["r"] + "}id"


@dataclass(frozen=True)
class ExtractionConfig:
    include_text: bool = True
    include_notes: bool = True
    include_workbooks: bool = True
    include_images: bool = True
    include_hidden_slides: bool = True
    include_hidden_sheets: bool = True
    include_unreferenced_embeddings: bool = True
    max_file_bytes: int = 256 * 1024 * 1024
    max_part_bytes: int = 128 * 1024 * 1024
    max_uncompressed_bytes: int = 512 * 1024 * 1024
    max_zip_entries: int = 20000
    max_sheet_cells: int = 250000
    max_chart_points: int = 250000


def local_path(path: str | Path) -> Path:
    """UC Volume paths work with ordinary Python; legacy DBFS needs FUSE."""
    value = str(path)
    if value.startswith("dbfs:/Volumes/"):
        value = value[len("dbfs:"):]
    elif value.startswith("dbfs:/"):
        value = "/dbfs/" + value[len("dbfs:/"):]
    elif value.startswith("file:/"):
        value = value[len("file:"):]
    elif re.match(r"^[A-Za-z][A-Za-z0-9+.-]*://", value):
        raise ValueError("Pass local/Volume paths or bytes; stage cloud URIs first.")
    return Path(value)


def _tag(node) -> str:
    return node.tag.rsplit("}", 1)[-1]


def _text(node) -> str:
    if node is None:
        return ""
    paragraphs = []
    for paragraph in node.findall(".//a:p", NS):
        chunks = []
        for child in paragraph.iter():
            if child.tag == "{" + NS["a"] + "}t":
                chunks.append(child.text or "")
            elif child.tag == "{" + NS["a"] + "}br":
                chunks.append("\n")
            elif child.tag == "{" + NS["a"] + "}tab":
                chunks.append("\t")
        paragraphs.append("".join(chunks))
    return "\n".join(paragraphs).strip()


def _value(node, path: str, default=None):
    found = node.find(path, NS)
    return found.get("val", default) if found is not None else default


def _json_value(value):
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return {"seconds": value.total_seconds(), "type": "timedelta"}
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    return str(value)


def _number(raw: str | None):
    if raw is None or raw == "":
        return None
    try:
        if re.fullmatch(r"[-+]?\d+", raw):
            return int(raw)
        result = float(raw)
        return result if math.isfinite(result) else raw
    except ValueError:
        return raw


def _checked_zip(blob: bytes, config: ExtractionConfig):
    archive = zipfile.ZipFile(io.BytesIO(blob))
    try:
        infos = archive.infolist()
        names = [item.filename for item in infos]
        if len(infos) > config.max_zip_entries or len(set(names)) != len(names):
            raise ValueError("Archive has too many entries or duplicate part names.")
        if sum(item.file_size for item in infos) > config.max_uncompressed_bytes:
            raise ValueError("Archive exceeds max_uncompressed_bytes.")
        if any(item.file_size > config.max_part_bytes for item in infos):
            raise ValueError("Archive part exceeds max_part_bytes.")
        return archive
    except Exception:
        archive.close()
        raise


def _xml(archive, part):

    return ET.fromstring(archive.read(part), forbid_dtd=True)


def _rels(archive, part):
    path = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if path not in archive.namelist():
        return {}
    result = {}
    for rel in _xml(archive, path):
        target = rel.get("Target", "")
        external = rel.get("TargetMode") == "External"
        resolved = target
        if not external:
            target = unquote(target)
            resolved = posixpath.normpath(
                target.lstrip("/") if target.startswith("/")
                else posixpath.join(posixpath.dirname(part), target)
            )
            if resolved == ".." or resolved.startswith("../") or "\\" in resolved:
                raise ValueError("Invalid internal relationship target.")
        result[rel.get("Id")] = {
            "target": resolved, "type": rel.get("Type", "").rsplit("/", 1)[-1],
            "external": external,
        }
    return result


def _shapes(root):
    """Visit grouped shapes and one supported AlternateContent branch only."""
    for child in root:
        name = _tag(child)
        if name == "AlternateContent":
            choices = child.findall("mc:Choice", NS)
            chosen = next((c for c in choices if c.find(".//c:chart", NS) is not None), None)
            if chosen is None:
                chosen = next(iter(choices), None)
            if chosen is None:
                chosen = child.find("mc:Fallback", NS)
            if chosen is not None:
                yield from _shapes(chosen)
        elif name == "grpSp":
            yield from _shapes(child)
        elif name in {"sp", "pic", "graphicFrame", "cxnSp", "contentPart"}:
            yield child


def _shape_info(shape):
    nv = shape.find(".//p:cNvPr", NS)
    ph = shape.find(".//p:ph", NS)
    return {
        "shape_id": nv.get("id") if nv is not None else None,
        "shape_name": nv.get("name", "") if nv is not None else "",
        "alt_text": nv.get("descr", "") if nv is not None else "",
        "alt_title": nv.get("title", "") if nv is not None else "",
        "placeholder_type": ph.get("type", "obj") if ph is not None else None,
    }


def _shape_layouts(root, slide_width, slide_height):
    """Map shapes to slide coordinates, including unrotated/scaled nested groups.

    Rotated/flipped groups are deliberately left unpositioned. Their data is
    still extracted, but a nearby label will never be assigned using bad geometry.
    """
    layouts = {}

    def walk(node, transform=(1.0, 1.0, 0.0, 0.0), reliable=True):
        sx, sy, tx, ty = transform
        for child in node:
            name = _tag(child)
            if name in {"AlternateContent", "Choice", "Fallback"}:
                walk(child, transform, reliable)
                continue
            if name == "grpSp":
                xfrm = child.find("p:grpSpPr/a:xfrm", NS)
                usable, next_transform = reliable, transform
                if xfrm is None:
                    usable = False
                else:
                    off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
                    ch_off, ch_ext = xfrm.find("a:chOff", NS), xfrm.find("a:chExt", NS)
                    usable = usable and all(n is not None for n in (off, ext, ch_off, ch_ext))
                    usable = usable and xfrm.get("rot", "0") == "0" and not any(
                        xfrm.get(k) in {"1", "true"} for k in ("flipH", "flipV"))
                    if usable and float(ch_ext.get("cx", 0)) and float(ch_ext.get("cy", 0)):
                        gx = float(ext.get("cx")) / float(ch_ext.get("cx"))
                        gy = float(ext.get("cy")) / float(ch_ext.get("cy"))
                        next_transform = (sx * gx, sy * gy,
                                          tx + sx * (float(off.get("x")) - gx * float(ch_off.get("x"))),
                                          ty + sy * (float(off.get("y")) - gy * float(ch_off.get("y"))))
                    else:
                        usable = False
                walk(child, next_transform, usable)
                continue
            if name not in {"sp", "pic", "graphicFrame", "cxnSp", "contentPart"}:
                continue
            xfrm = child.find("p:xfrm", NS)
            if xfrm is None:
                xfrm = child.find("p:spPr/a:xfrm", NS)
            box, position = None, "unknown"
            geometry_note = None
            if reliable and xfrm is not None and xfrm.get("rot", "0") == "0":
                off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
                if off is not None and ext is not None:
                    x, y = tx + sx * float(off.get("x")), ty + sy * float(off.get("y"))
                    width, height = sx * float(ext.get("cx")), sy * float(ext.get("cy"))
                    box = {"x": x, "y": y, "width": width, "height": height, "unit": "EMU"}
                    if slide_width and slide_height:
                        center_x, center_y = (x + width / 2) / slide_width, (y + height / 2) / slide_height
                        horizontal = "left" if center_x < 0.4 else "right" if center_x > 0.6 else "center"
                        vertical = "top" if center_y < 0.4 else "bottom" if center_y > 0.6 else "middle"
                        position = f"{vertical}-{horizontal}"
            if box is None:
                geometry_note = "Shape position unavailable (missing, inherited, or rotated/group-transformed geometry)."
            layouts[id(child)] = {"bounding_box": box, "position_on_slide": position, "geometry_note": geometry_note}

    if root is not None:
        walk(root)
    return layouts


def _nearby_labels(shape, shapes, layouts):
    """Return candidate text above a shape, without treating proximity as proof."""
    box = layouts.get(id(shape), {}).get("bounding_box")
    if not box or box["width"] <= 0:
        return []
    candidates = []
    for other in shapes:
        info = _shape_info(other)
        if _tag(other) != "sp" or info["placeholder_type"] in {"title", "ctrTitle", "dt", "ftr", "sldNum"}:
            continue
        text = _text(other)
        obox = layouts.get(id(other), {}).get("bounding_box")
        if not text or not obox or obox["width"] <= 0:
            continue
        gap = box["y"] - (obox["y"] + obox["height"])
        overlap = max(0.0, min(box["x"] + box["width"], obox["x"] + obox["width"]) - max(box["x"], obox["x"]))
        coverage = overlap / min(box["width"], obox["width"])

        if 0 <= gap <= 0.8 * 914400 and coverage >= 0.5:
            candidates.append({"text": text, "shape_id": info["shape_id"], "bounding_box": obox,
                               "vertical_gap_emu": gap, "horizontal_overlap_fraction": coverage,
                               "association": "unverified_proximity_candidate"})
    return sorted(candidates, key=lambda c: (c["vertical_gap_emu"], -c["horizontal_overlap_fraction"]))[:3]


def _assess_record(record):
    """Flag observable problems, not a calibrated probability of correctness."""
    flags = []

    def flag(code, dimension, detail):
        if not any(f["code"] == code for f in flags):
            flags.append({"code": code, "dimension": dimension, "detail": detail})

    kind, data = record["kind"], record["data"]
    evidence = record["description_evidence"]
    extraction = "no_detected_issue"
    meaning = "not_assessed"
    freshness = "not_assessed"
    if record["status"] in {"needs_review", "partial", "reference_only"}:
        extraction = "partial" if record["status"] == "partial" else "unavailable_or_not_decoded"
        flag("INCOMPLETE_OR_UNSUPPORTED_CONTENT", "extraction", "; ".join(record["warnings"]) or "Content was not fully extracted.")
    if kind == "chart":
        meaning, freshness = "source_labeled", "unverified_saved_values"
        local_label = data.get("chart_title") or evidence.get("alt_title") or evidence.get("alt_text")
        if not local_label:
            if record.get("chart_count_on_slide", 0) > 1:
                flag("MULTIPLE_CHARTS_WITHOUT_LOCAL_LABEL", "meaning",
                     "This chart has no intrinsic title/alternative label. The slide title alone cannot establish this chart's meaning.")
            elif not record.get("slide_title"):
                flag("MISSING_CHART_CONTEXT", "meaning", "Neither chart nor slide has a title/alternative label.")
            if evidence.get("nearby_text_candidates"):
                flag("NEARBY_LABEL_UNVERIFIED", "meaning", "Nearby text is a candidate label; spatial proximity does not prove it belongs to this chart.")
        if any(s["series_name"].endswith("(unlabeled)") for s in data.get("series", [])):
            flag("MISSING_SERIES_LABEL", "meaning", "One or more series has no saved label.")
        if record.get("chart_count_on_slide", 0) > 1 and not record.get("bounding_box"):
            flag("CHART_POSITION_UNRESOLVED", "association", "Data stays separate by chart ID, but its position on the slide could not be resolved.")
        if record.get("duplicate_chart_labels_on_slide"):
            flag("DUPLICATE_CHART_LABELS", "meaning", "Multiple charts on this slide have identical titles/series/axis labels. IDs and positions distinguish the objects, but the labels do not establish their different meanings.")
        for series in data.get("series", []):
            sources = series["sources"]
            if sources["val"]["present"] and not sources["cat"]["present"]:
                flag("MISSING_CATEGORY_LABELS", "meaning", "A categorical series has values but no category vector.")
            active = [key for key in ("cat", "val", "xVal", "yVal", "bubbleSize") if sources[key]["present"]]
            counts = {sources[key]["declared_count"] for key in active if sources[key]["declared_count"] is not None}
            if len(counts) > 1:
                flag("MISMATCHED_POINT_COUNTS", "extraction", "Saved category/value/x/y/size vectors have different declared lengths; indices were retained.")
            fields = {"cat": "category", "val": "value", "xVal": "x", "yVal": "y", "bubbleSize": "bubble_size"}
            for row in (r for r in data.get("rows", []) if r["series_number"] == series["series_number"]):
                for key in active:
                    value = row.get(fields[key])
                    if value is None:
                        flag("MISSING_CHART_POINTS", "extraction", "Some active chart-vector positions have no saved values. These may be intentional blanks; no values were invented.")
                    elif key != "cat" and not isinstance(value, (int, float)):
                        flag("NONNUMERIC_CHART_VALUE", "extraction", "A numeric chart vector contains an error token or nonnumeric value; the original text is retained.")
        if not data.get("rows"):
            flag("NO_CHART_DATA", "extraction", "No chart data rows could be recovered.")
    elif kind == "embedded_worksheet":
        meaning = "source_labels_only"
        if data.get("formulas"):
            freshness = "unverified_saved_values"
        if any(f.get("cached_value") is None for f in data.get("formulas", [])):
            flag("MISSING_FORMULA_RESULT", "extraction", "One or more formulas lacks a saved result; the formula is preserved and its result remains null.")
        if any(m.get("data_type") == "e" for m in data.get("cell_metadata", [])):
            flag("WORKSHEET_ERROR_VALUE", "extraction", "The worksheet contains an Excel error value.")
        if not data.get("source_references"):
            flag("UNLINKED_WORKBOOK", "association", "No selected slide/chart reference establishes what this embedded workbook represents.")
        if data.get("rows") and not data.get("label_candidates"):
            flag("WORKSHEET_WITHOUT_LABELS", "meaning", "No text labels were found in the stored worksheet range.")
    elif kind == "table":
        meaning = "source_labels_only"
        if not data.get("first_row_styled_as_header"):
            flag("TABLE_HEADER_UNCONFIRMED", "meaning", "There is no first-row header styling. All rows are retained; a header row was not assumed.")
        if not record.get("slide_title") and not evidence.get("alt_text") and not evidence.get("alt_title"):
            flag("TABLE_CONTEXT_MISSING", "meaning", "No slide or alternative label establishes the table's subject.")
    if record.get("description_method") == "user_callback_inferred":
        meaning = "inferred"
        flag("INFERRED_DESCRIPTION", "meaning", "The optional callback produced this description; its semantic correctness has not been verified.")
    elif any(f["dimension"] == "meaning" for f in flags):
        meaning = "ambiguous"
    if any(f["dimension"] == "extraction" for f in flags) and extraction == "no_detected_issue":
        extraction = "check_required"
    record["is_uncertain"] = bool(flags)
    record["needs_review"] = bool(flags)
    record["uncertainty_reasons"] = flags
    record["quality"] = {"assessment_method": "deterministic_rules_not_calibrated_probability",
                         "data_extraction": extraction, "meaning": meaning, "source_freshness": freshness,
                         "note": "No flag means no listed rule detected a problem; it is not a guarantee of completeness or correctness."}


def _refresh_review_counts(report):
    report["review_record_count"] = sum(r["needs_review"] for r in report["records"])
    report["uncertain_record_count"] = sum(r["is_uncertain"] for r in report["records"])


def _chart_vector(node, config):
    result = {"formula": None, "number_format": None, "declared_count": None,
              "points": {}, "levels": [], "warnings": [], "present": node is not None}
    if node is None:
        return result
    f = node.find(".//c:f", NS)
    result["formula"] = f.text if f is not None else None
    cache = next((n for n in node.iter() if _tag(n) in
                  {"numCache", "strCache", "numLit", "strLit", "multiLvlStrCache"}), None)
    if cache is None:
        literal = node.find("c:v", NS)
        if literal is not None:
            result["points"][0] = {"value": literal.text, "raw": literal.text}
        elif result["formula"]:
            result["warnings"].append("Reference has no saved cache; values were not inferred from its formula.")
        return result
    count = _value(cache, "c:ptCount")
    result["declared_count"] = int(count) if count is not None else None
    fmt = cache.find("c:formatCode", NS)
    result["number_format"] = fmt.text if fmt is not None else None
    numeric = _tag(cache) in {"numCache", "numLit"}
    containers = cache.findall("c:lvl", NS) if _tag(cache) == "multiLvlStrCache" else [cache]
    if (result["declared_count"] or 0) > config.max_chart_points:
        raise ValueError("Chart exceeds max_chart_points.")
    for container in containers:
        points = {}
        for point in container.findall("c:pt", NS):
            index = int(point.get("idx"))
            if index < 0 or index >= config.max_chart_points:
                raise ValueError("Chart point index exceeds limits.")
            if index in points:
                result["warnings"].append(f"Duplicate point index {index}; last saved entry retained.")
            val = point.find("c:v", NS)
            raw = val.text if val is not None else None
            points[index] = {"value": _number(raw) if numeric else raw, "raw": raw}
        if _tag(cache) == "multiLvlStrCache":
            result["levels"].append(points)
        else:
            result["points"] = points
    if result["levels"]:
        result["points"] = result["levels"][0]
    return result


def _chart_data(root, config):
    chart = root.find("c:chart", NS)
    if chart is None:
        raise ValueError("Unsupported chart namespace/format (e.g. chartEx).")
    title_node = chart.find("c:title", NS)
    title = _text(title_node)
    if not title and title_node is not None:
        title = " ".join(n.text or "" for n in title_node.findall(".//c:v", NS))
    plot = chart.find("c:plotArea", NS)
    if plot is None:
        raise ValueError("Chart has no plotArea.")
    axes = [{"axis_type": _tag(ax), "axis_id": _value(ax, "c:axId"),
             "title": _text(ax.find("c:title", NS)),
             "position": _value(ax, "c:axPos"),
             "number_format": (ax.find("c:numFmt", NS).get("formatCode")
                               if ax.find("c:numFmt", NS) is not None else None)}
            for ax in plot if _tag(ax).endswith("Ax")]
    series_out, rows, warnings = [], [], [
        "Chart values are saved caches/literals and may be stale relative to the workbook. Raw numeric text is retained."
    ]
    chart_types = []
    for group in plot:
        if not _tag(group).endswith("Chart"):
            continue
        chart_type = _tag(group)
        chart_types.append(chart_type)
        for ser in group.findall("c:ser", NS):
            vectors = {key: _chart_vector(ser.find("c:" + key, NS), config)
                       for key in ("tx", "cat", "val", "xVal", "yVal", "bubbleSize")}
            series_number = len(series_out) + 1
            name = next((p["value"] for p in vectors["tx"]["points"].values()
                         if p["value"] is not None), None)
            name = str(name) if name is not None else f"Series {series_number} (unlabeled)"
            series_warnings = [w for vec in vectors.values() for w in vec["warnings"]]
            indices = set()
            for key, vec in vectors.items():
                if key == "tx":
                    continue
                indices.update(vec["points"])
                for level in vec["levels"]:
                    indices.update(level)
                indices.update(range(vec["declared_count"] or 0))
            if len(rows) + len(indices) > config.max_chart_points:
                raise ValueError("Combined chart rows exceed max_chart_points.")
            series_out.append({
                "series_number": series_number, "series_name": name,
                "series_index": _value(ser, "c:idx"), "series_order": _value(ser, "c:order"),
                "chart_type": chart_type,
                "axis_ids": [n.get("val") for n in group.findall("c:axId", NS)],
                "sources": {key: {k: v for k, v in vec.items() if k not in {"points", "levels", "warnings"}}
                            for key, vec in vectors.items()},
                "warnings": series_warnings,
            })
            warnings.extend(series_warnings)
            if not indices:
                warnings.append(f"{name}: no cached data points available.")
            for index in sorted(indices):
                row = {"series_number": series_number, "series_name": name,
                       "chart_type": chart_type, "point_index": index}
                for key, field in {"cat": "category", "val": "value", "xVal": "x",
                                   "yVal": "y", "bubbleSize": "bubble_size"}.items():
                    point = vectors[key]["points"].get(index, {})
                    row[field] = point.get("value")
                    row[field + "_raw"] = point.get("raw")
                row["category_levels"] = [level.get(index, {}).get("value") for level in vectors["cat"]["levels"]]
                rows.append(row)
    if not series_out:
        warnings.append("No supported chart series found; this chart needs review.")
    return {"chart_title": title, "chart_types": chart_types, "axes": axes,
            "series": series_out, "rows": rows,
            "category_level_order": "XML order (innermost first); omitted hierarchy labels remain null.",
            "date_system": "1904" if _value(root, "c:date1904", "0") in {"1", "true"} else "1900",
            "plot_visible_only": _value(chart, "c:plotVisOnly"),
            "display_blanks_as": _value(chart, "c:dispBlanksAs")}, list(dict.fromkeys(warnings))


def _table_data(table):
    rows, cell_metadata = [], []
    for r, tr in enumerate(table.findall("a:tr", NS), 1):
        row = []
        for c, cell in enumerate(tr.findall("a:tc", NS), 1):
            row.append(_text(cell))
            if cell.attrib:
                cell_metadata.append({"row": r, "column": c, **cell.attrib})
        rows.append(row)
    prop = table.find("a:tblPr", NS)
    first_row = prop is not None and prop.get("firstRow") in {"1", "true"}
    return {"columns": [f"column_{i+1}" for i in range(max(map(len, rows), default=0))],
            "rows": rows, "first_row_text": rows[0] if rows else [],
            "first_row_styled_as_header": first_row, "merge_metadata": cell_metadata,
            "value_type": "text_as_stored; percentages/currency are not coerced"}


def _worksheet_data(archive, sheet_root, ws, cached_ws, config):
    cells = sheet_root.findall("x:sheetData/x:row/x:c", NS)
    stored = []
    for node in cells:
        if node.find("x:v", NS) is not None or node.find("x:f", NS) is not None or node.find("x:is", NS) is not None:
            stored.append(node)
    if not stored:
        return {"range": None, "columns": [], "rows": [], "formulas": [], "cell_metadata": [],
                "label_candidates": [], "merged_ranges": [], "hidden_rows": [], "hidden_column_ranges": []}, []
    coords = [coordinate_to_tuple(node.get("r")) for node in stored]
    min_r, max_r = min(r for r, _ in coords), max(r for r, _ in coords)
    min_c, max_c = min(c for _, c in coords), max(c for _, c in coords)
    if (max_r - min_r + 1) * (max_c - min_c + 1) > config.max_sheet_cells:
        raise ValueError("Worksheet bounding rectangle exceeds max_sheet_cells; raise the explicit limit if appropriate.")

    ws.reset_dimensions()
    cached_ws.reset_dimensions()
    kwargs = dict(min_row=min_r, max_row=max_r, min_col=min_c, max_col=max_c)
    rows, formulas, metadata, labels = [], [], [], []
    raw_by_coordinate = {node.get("r"): node.findtext("x:v", namespaces=NS) for node in stored}
    missing_cache = 0
    for source_row, value_row in zip(ws.iter_rows(**kwargs), cached_ws.iter_rows(**kwargs)):
        out = []
        for cell, cached in zip(source_row, value_row):
            value = cached.value if cell.data_type == "f" else cell.value
            out.append(_json_value(value))
            if cell.value is None:
                continue
            coordinate = cell.coordinate
            if cell.data_type == "f":
                formula = cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", str(cell.value))
                formulas.append({"cell": coordinate, "formula": formula, "cached_value": _json_value(cached.value),
                                 "cached_raw": raw_by_coordinate.get(coordinate)})
                missing_cache += cached.value is None
            if cell.data_type != "s" or cell.number_format != "General":
                metadata.append({"cell": coordinate, "data_type": cell.data_type,
                                 "number_format": cell.number_format, "raw_xml_value": raw_by_coordinate.get(coordinate)})
            if isinstance(value, str) and cell.data_type != "f" and len(labels) < 12:
                labels.append({"cell": coordinate, "text": value})
        rows.append(out)
    warnings = []
    if formulas:
        warnings.append("Formula results are saved values, not recalculated; the original formulas are retained.")
    if missing_cache:
        warnings.append(f"{missing_cache} formula cell(s) lack cached results; their values remain null.")
    merged = [n.get("ref") for n in sheet_root.findall("x:mergeCells/x:mergeCell", NS)]
    data = {"range": f"{get_column_letter(min_c)}{min_r}:{get_column_letter(max_c)}{max_r}",
            "first_excel_row": min_r,
            "columns": [get_column_letter(c) for c in range(min_c, max_c + 1)],
            "rows": rows, "formulas": formulas, "cell_metadata": metadata,
            "label_candidates": labels, "merged_ranges": merged,
            "hidden_rows": [int(n.get("r")) for n in sheet_root.findall("x:sheetData/x:row", NS)
                            if n.get("hidden") in {"1", "true"}],
            "hidden_column_ranges": [{"min": int(n.get("min")), "max": int(n.get("max"))}
                                     for n in sheet_root.findall("x:cols/x:col", NS)
                                     if n.get("hidden") in {"1", "true"}]}
    return data, warnings


class _Extractor:
    def __init__(self, archive, report, config):
        self.z, self.report, self.config = archive, report, config
        self.embedding_sources = {}
        self.names = set(archive.namelist())

    def record(self, kind, context, part, description, data, warnings=None, status="extracted", **extra):
        sequence = len(self.report["records"])
        identity = f"{self.report['source_sha256']}|{context.get('slide_number')}|{part}|{kind}|{sequence}"
        record = {
            "record_id": hashlib.sha256(identity.encode()).hexdigest()[:24],
            "source_file": self.report["source_file"], "source_sha256": self.report["source_sha256"],
            "slide_number": context.get("slide_number"), "slide_title": context.get("slide_title", ""),
            "hidden_slide": context.get("hidden_slide"), "slide_part": context.get("slide_part"),
            "kind": kind, "source_part": part, "shape_id": context.get("shape_id"),
            "shape_name": context.get("shape_name", ""),
            "chart_index_on_slide": context.get("chart_index_on_slide"),
            "chart_count_on_slide": context.get("chart_count_on_slide"),
            "bounding_box": context.get("bounding_box"),
            "position_on_slide": context.get("position_on_slide", "unknown"),
            "what_data_represents": description,
            "description_method": "source_labels_and_context",
            "description_evidence": {k: context.get(k) for k in
                                     ("slide_title", "slide_text", "alt_text", "alt_title", "nearby_text_candidates", "geometry_note") if context.get(k)},
            "status": status, "warnings": warnings or [], "data": data,
            **extra,
        }
        if context.get("chart_count_on_slide", 0) > 1 and context.get("slide_text"):
            record["description_evidence"]["slide_text_scope"] = "Shared slide context; not necessarily specific to this chart/object."
        _assess_record(record)
        self.report["records"].append(record)
        return record

    def issue(self, context, part, message, kind="extraction_issue", data=None):
        return self.record(kind, context, part, f"Content needing review: {message}", data or {}, [message], "needs_review")

    def register_embeddings(self, part, context):
        for rid, rel in _rels(self.z, part).items():
            if rel["external"]:
                self.record("external_link", context, part, f"External {rel['type']} reference", rel,
                            ["External content was not fetched or refreshed."], "reference_only", relationship_id=rid)
            elif rel["target"].startswith("ppt/embeddings/") or rel["type"] in {"package", "oleObject"}:
                entry = {k: context.get(k) for k in ("slide_number", "slide_title", "hidden_slide", "slide_part", "shape_id", "shape_name", "chart_index_on_slide", "chart_count_on_slide", "bounding_box", "position_on_slide")}
                entry["referencing_part"] = part
                entry["relationship_id"] = rid
                sources = self.embedding_sources.setdefault(rel["target"], [])
                if entry not in sources:
                    sources.append(entry)

    def slide(self, part, number):
        root = _xml(self.z, part)
        hidden = root.get("show") in {"0", "false"}
        if hidden and not self.config.include_hidden_slides:
            self.report["warnings"].append(f"Hidden slide {number} skipped by configuration.")
            return
        tree = root.find("p:cSld/p:spTree", NS)
        shapes = list(_shapes(tree)) if tree is not None else []
        slide_size = self.report.get("slide_size_emu", {})
        try:
            layouts = _shape_layouts(tree, slide_size.get("width"), slide_size.get("height"))
        except (TypeError, ValueError, KeyError) as exc:


            layouts = {id(s): {"bounding_box": None, "position_on_slide": "unknown",
                               "geometry_note": f"Invalid shape geometry: {exc}"} for s in shapes}
        chart_refs = [ref for shape in shapes for ref in shape.iter()
                      if _tag(ref) == "chart" and ref.get(RID)]
        chart_indices = {id(ref): index for index, ref in enumerate(chart_refs, 1)}
        title = next((_text(s) for s in shapes if _shape_info(s)["placeholder_type"] in {"title", "ctrTitle"}), "")
        slide_text = "\n".join(_text(s) for s in shapes if _tag(s) == "sp" and _text(s))
        context = {"slide_number": number, "slide_title": title, "slide_part": part,
                   "slide_text": slide_text, "hidden_slide": hidden,
                   "chart_count_on_slide": len(chart_refs)}
        relationships = _rels(self.z, part)
        if self.config.include_text:
            self.record("slide_text", context, part, title or f"Text on slide {number}", {"text": slide_text})
        self.register_embeddings(part, context)
        for shape in shapes:
            shape_context = {**context, **_shape_info(shape), **layouts.get(id(shape), {})}
            shape_context["nearby_text_candidates"] = _nearby_labels(shape, shapes, layouts)
            ctx = shape_context
            label = title or f"Slide {number}"
            table = shape.find(".//a:tbl", NS)
            if table is not None:
                data = _table_data(table)
                header = ", ".join(t for t in data["first_row_text"] if t)[:250]
                description = f"Table on '{label}'" + (f"; first row: {header}" if header else "")
                self.record("table", ctx, part, description, data)
            for chart_ref in shape.findall(".//c:chart", NS):
                ctx = {**shape_context, "chart_index_on_slide": chart_indices.get(id(chart_ref))}
                rel = relationships.get(chart_ref.get(RID))
                if not rel or rel["external"]:
                    self.issue(ctx, part, "Chart relationship is missing or external.")
                    continue
                chart_part = rel["target"]
                try:
                    self.register_embeddings(chart_part, ctx)
                    data, warnings = _chart_data(_xml(self.z, chart_part), self.config)
                    names = ", ".join(s["series_name"] for s in data["series"][:8])
                    axes = ", ".join(ax["title"] for ax in data["axes"] if ax["title"])
                    own_label = data["chart_title"] or ctx["alt_title"] or ctx["alt_text"]
                    chart_identity = f"Chart {ctx['chart_index_on_slide']} of {len(chart_refs)}"
                    if ctx.get("position_on_slide") not in {None, "unknown"}:
                        chart_identity += f" ({ctx['position_on_slide']})"
                    description = f"{chart_identity}: '{own_label or label}' on slide {number}"
                    if not own_label and ctx["nearby_text_candidates"]:
                        description += f"; nearby label candidate (unverified): {ctx['nearby_text_candidates'][0]['text']}"
                    if names:
                        description += f"; series: {names}"
                    if axes:
                        description += f"; axes: {axes}"
                    status = "saved_values" if data["rows"] else "needs_review"
                    if any(s["warnings"] for s in data["series"]) or any("no cached data" in w for w in warnings):
                        status = "partial"
                    self.record("chart", ctx, chart_part, description, data, warnings, status)
                except Exception as exc:
                    self.issue(ctx, chart_part, f"Chart could not be fully extracted: {type(exc).__name__}: {exc}")

            for ref in shape.iter():
                if _tag(ref) == "chart" and ref.tag != "{" + NS["c"] + "}chart":
                    ctx = {**shape_context, "chart_index_on_slide": chart_indices.get(id(ref))}
                    rel = relationships.get(ref.get(RID))
                    modern_part = rel["target"] if rel and not rel["external"] else part
                    try:
                        if modern_part != part:
                            self.register_embeddings(modern_part, ctx)
                    except Exception as exc:
                        self.report["warnings"].append(f"{modern_part}: {exc}")
                    self.issue(ctx, modern_part, "Modern chartEx chart requires a format-specific parser; linked embedded workbooks are extracted separately.")
            ctx = shape_context
            for dgm in shape.findall(".//dgm:relIds", NS):
                rid = dgm.get("{" + NS["r"] + "}dm")
                rel = relationships.get(rid)
                if rel and not rel["external"]:
                    try:
                        dpart = rel["target"]
                        self.record("diagram_text", ctx, dpart, f"SmartArt labels on '{label}'",
                                    {"text": _text(_xml(self.z, dpart))},
                                    ["Labels only; diagram geometry and relationships are not interpreted."], "partial")
                    except Exception as exc:
                        self.issue(ctx, part, f"SmartArt extraction failed: {exc}")
            if self.config.include_images:
                for image_index, blip in enumerate(shape.findall(".//a:blip", NS), 1):
                    rid = blip.get("{" + NS["r"] + "}embed") or blip.get("{" + NS["r"] + "}link")
                    rel = relationships.get(rid)
                    if rel:
                        self.record("image_reference", ctx, rel["target"],
                                    f"Image on '{label}'" + (f"; alt text: {ctx['alt_text']}" if ctx["alt_text"] else ""),
                                    {"media_part": rel["target"], "external": rel["external"], "image_index": image_index},
                                    ["Image inventory only. Any numbers in this image require OCR/vision and validation."], "reference_only")


        label_groups = {}
        for record in self.report["records"]:
            if record["kind"] != "chart" or record["slide_number"] != number:
                continue
            data = record["data"]
            signature = (data["chart_title"], record["description_evidence"].get("alt_title"),
                         record["description_evidence"].get("alt_text"),
                         tuple(s["series_name"] for s in data["series"]),
                         tuple(ax["title"] for ax in data["axes"]))
            label_groups.setdefault(signature, []).append(record)
        for group in label_groups.values():
            if len(group) > 1:
                for record in group:
                    record["duplicate_chart_labels_on_slide"] = True
                    _assess_record(record)
        if self.config.include_notes:
            for rel in relationships.values():
                if rel["type"] == "notesSlide" and not rel["external"]:
                    try:
                        note_root = _xml(self.z, rel["target"])
                        note_shapes = note_root.findall(".//p:sp", NS)
                        notes = "\n".join(_text(s) for s in note_shapes
                                          if _shape_info(s)["placeholder_type"] not in {"sldImg", "sldNum", "dt", "hdr", "ftr"} and _text(s))
                        if notes:
                            self.record("speaker_notes", context, rel["target"], f"Speaker notes for '{title or 'Slide ' + str(number)}'", {"text": notes})
                    except Exception as exc:
                        self.issue(context, rel["target"], f"Notes extraction failed: {exc}")

    def workbooks(self):
        parts = set(self.embedding_sources)
        if self.config.include_unreferenced_embeddings:
            parts.update(p for p in self.names if p.startswith("ppt/embeddings/") and not p.endswith("/"))
        for part in sorted(parts):
            sources = self.embedding_sources.get(part, [])

            ctx = dict(sources[0]) if len({s["slide_number"] for s in sources}) == 1 else {}
            if len({(s.get("referencing_part"), s.get("shape_id")) for s in sources}) > 1:
                for key in ("shape_id", "shape_name", "chart_index_on_slide", "bounding_box", "position_on_slide"):
                    ctx.pop(key, None)
            if not sources:
                ctx = {}
            try:
                blob = self.z.read(part)
                if not zipfile.is_zipfile(io.BytesIO(blob)):
                    self.issue(ctx, part, "Legacy OLE/.xls/.bin or other embedded binary is not decoded.",
                               "unsupported_embedding", {"byte_count": len(blob), "source_references": sources})
                    continue
                with _checked_zip(blob, self.config) as embedded:
                    if "xl/workbook.xml" not in embedded.namelist():
                        self.issue(ctx, part, "Embedded package is not an OOXML workbook.",
                                   "unsupported_embedding", {"source_references": sources})
                        continue
                    workbook_root = _xml(embedded, "xl/workbook.xml")
                    relationships = _rels(embedded, "xl/workbook.xml")
                    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=False, keep_links=False)
                    values = None
                    try:
                        values = load_workbook(io.BytesIO(blob), read_only=True, data_only=True, keep_links=False)
                        for sheet in workbook_root.findall("x:sheets/x:sheet", NS):
                            name = sheet.get("name")
                            state = sheet.get("state", "visible")
                            if state != "visible" and not self.config.include_hidden_sheets:
                                self.report["warnings"].append(f"Hidden sheet '{name}' in {part} skipped by configuration.")
                                continue
                            sheet_rel = relationships.get(sheet.get(RID))
                            try:
                                if not sheet_rel or sheet_rel["type"] != "worksheet" or sheet_rel["external"]:
                                    raise ValueError("Unsupported workbook sheet type.")
                                data, warnings = _worksheet_data(embedded, _xml(embedded, sheet_rel["target"]), wb[name], values[name], self.config)
                                data.update({"sheet_name": name, "sheet_state": state,
                                             "excel_epoch": wb.epoch.isoformat(), "source_references": sources})
                                labels = ", ".join(item["text"] for item in data["label_candidates"][:6])[:250]
                                slide_names = ", ".join(str(n) for n in sorted({s["slide_number"] for s in sources if s["slide_number"] is not None}))
                                description = f"Embedded worksheet '{name}'"
                                description += f" referenced by slide(s) {slide_names}" if slide_names else " (no selected-slide reference found)"
                                if labels:
                                    description += f"; label examples: {labels}"
                                warnings.insert(0, "Full stored worksheet range; may include data not plotted or visible on the slide. Header semantics are not assumed.")
                                self.record("embedded_worksheet", ctx, part, description, data, warnings,
                                            "saved_values" if data["formulas"] else "extracted", sheet_name=name)
                            except Exception as exc:
                                self.issue(ctx, part, f"Worksheet '{name}' extraction failed: {type(exc).__name__}: {exc}",
                                           data={"sheet_name": name, "source_references": sources})
                    finally:
                        wb.close()
                        if values is not None:
                            values.close()
            except Exception as exc:
                self.issue(ctx, part, f"Embedded object extraction failed: {type(exc).__name__}: {exc}", data={"source_references": sources})


def extract_pptx(source: str | Path | bytes | bytearray, *, source_name: str | None = None,
                 config: ExtractionConfig | None = None) -> dict:
    """Return a JSON-serializable report. Invalid packages raise; local object errors are recorded.

    Extraction is read-only. No external links, VBA, formulas, or AI calls execute.
    The whole deck is read on the current Python process (normally the driver).
    """
    config = config or ExtractionConfig()
    if isinstance(source, (bytes, bytearray)):
        blob = bytes(source)
        name = source_name or "presentation.pptx"
    else:
        path = local_path(source)
        if path.stat().st_size > config.max_file_bytes:
            raise ValueError("Presentation exceeds max_file_bytes.")
        blob = path.read_bytes()
        name = source_name or str(path)
    if len(blob) > config.max_file_bytes:
        raise ValueError("Presentation exceeds max_file_bytes.")
    report = {"schema_version": "1.1", "extractor_version": __version__, "source_file": name,
              "source_sha256": hashlib.sha256(blob).hexdigest(),
              "extracted_at_utc": datetime.now(timezone.utc).isoformat(),
              "slide_count": 0, "records": [], "warnings": [],
              "scope": "Slide text, native tables, classic chart caches, embedded OOXML worksheet values/formulas, notes, SmartArt labels and image references. No rendering/OCR, master/layout text, comments, or exhaustive metadata inventory."}
    with _checked_zip(blob, config) as archive:
        if "ppt/presentation.xml" not in archive.namelist():
            raise ValueError("Not an unencrypted PPTX/PPTM package. Convert legacy .ppt or encrypted files first.")
        root = _xml(archive, "ppt/presentation.xml")
        if root.tag != "{" + NS["p"] + "}presentation":
            raise ValueError("Unsupported presentation namespace (Strict OOXML is not currently supported).")
        relationships = _rels(archive, "ppt/presentation.xml")
        size = root.find("p:sldSz", NS)
        report["slide_size_emu"] = {"width": int(size.get("cx")), "height": int(size.get("cy"))} if size is not None else {}
        slides = root.findall("p:sldIdLst/p:sldId", NS)
        report["slide_count"] = len(slides)
        engine = _Extractor(archive, report, config)
        for number, slide in enumerate(slides, 1):
            rel = relationships.get(slide.get(RID))
            if not rel or rel["external"]:
                engine.issue({"slide_number": number}, "ppt/presentation.xml", "Missing/internal slide relationship could not be resolved.")
                continue
            try:
                engine.slide(rel["target"], number)
            except Exception as exc:
                engine.issue({"slide_number": number}, rel["target"], f"Slide extraction failed: {type(exc).__name__}: {exc}")
        if config.include_workbooks:
            engine.workbooks()
    report["record_count"] = len(report["records"])
    _refresh_review_counts(report)
    return report


def iter_pptx_reports(paths: Iterable[str | Path], *, config: ExtractionConfig | None = None):
    """Yield one result/error envelope per file, without retaining the whole batch."""
    for path in paths:
        try:
            report = extract_pptx(path, config=config)
            yield {"source_file": str(path), "ok": True, "report": report, "error": None}
        except Exception as exc:
            yield {"source_file": str(path), "ok": False, "report": None,
                   "error": f"{type(exc).__name__}: {exc}"}


def _dump(value, **kwargs):
    return json.dumps(value, ensure_ascii=False, allow_nan=False, **kwargs)


def write_outputs(report: dict, output_dir: str | Path) -> dict[str, str]:
    """Write a complete JSON report and JSONL records; a repeat run replaces these files."""
    directory = local_path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^A-Za-z0-9_.-]+", "_", Path(report["source_file"]).stem) or "presentation"
    stem += "_" + report["source_sha256"][:12]
    json_path, jsonl_path = directory / (stem + ".json"), directory / (stem + ".jsonl")
    json_path.write_text(_dump(report, indent=2), encoding="utf-8")
    with jsonl_path.open("w", encoding="utf-8") as stream:
        for record in report["records"]:
            stream.write(_dump(record) + "\n")
    return {"report_json": str(json_path), "records_jsonl": str(jsonl_path)}


def spark_schema():
    """Explicit schema avoids mixed Excel value types and empty/all-null inference errors."""
    from pyspark.sql.types import StructType, StructField, StringType, IntegerType, BooleanType
    string_columns = ["record_id", "source_file", "source_sha256", "kind", "slide_title",
                      "source_part", "shape_name", "what_data_represents", "description_method", "status",
                      "data_json", "evidence_json", "warnings_json", "record_json", "extracted_at_utc",
                      "position_on_slide", "bounding_box_json", "uncertainty_reasons_json", "quality_json"]
    return StructType([StructField(c, StringType(), True) for c in string_columns] +
                      [StructField("slide_number", IntegerType(), True), StructField("hidden_slide", BooleanType(), True),
                       StructField("chart_index_on_slide", IntegerType(), True), StructField("chart_count_on_slide", IntegerType(), True),
                       StructField("is_uncertain", BooleanType(), True), StructField("needs_review", BooleanType(), True)])


def to_spark_dataframe(spark, report: dict):
    """One row per dataset/object; heterogeneous nested payloads live in data_json.

    For large batches, persist and release each report rather than collecting all decks.
    """
    schema = spark_schema()
    rows = []
    for record in report["records"]:
        row = {field.name: record.get(field.name) for field in schema.fields}
        row.update(data_json=_dump(record["data"]), evidence_json=_dump(record["description_evidence"]),
                   warnings_json=_dump(record["warnings"]), record_json=_dump(record),
                   bounding_box_json=_dump(record.get("bounding_box")),
                   uncertainty_reasons_json=_dump(record["uncertainty_reasons"]), quality_json=_dump(record["quality"]),
                   extracted_at_utc=report["extracted_at_utc"])
        rows.append(row)
    return spark.createDataFrame(rows, schema=schema)


def enrich_descriptions(report: dict, describe: Callable[[dict], str]) -> dict:
    """Optional user-supplied semantic labeling; no model/provider is configured here.

    The callback receives a COPY of each dataset record (which may contain sensitive
    workbook data). It returns description text only. Numeric payloads stay unchanged.
    Any inferred description is marked separately from source-derived labels.
    """
    enriched = copy.deepcopy(report)
    for record in enriched["records"]:
        if record["kind"] not in {"chart", "table", "embedded_worksheet"}:
            continue
        description = describe(copy.deepcopy(record))
        if not isinstance(description, str) or not description.strip():
            raise ValueError("describe(record) must return a nonempty string.")
        record["source_label_description"] = record["what_data_represents"]
        record["what_data_represents"] = description.strip()
        record["description_method"] = "user_callback_inferred"
        _assess_record(record)
    _refresh_review_counts(enriched)
    return enriched
