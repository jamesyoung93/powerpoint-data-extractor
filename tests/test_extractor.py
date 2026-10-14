"""Run: python -m unittest discover -s tests -v

python-pptx is used ONLY to generate realistic regression fixtures.
"""
import copy
import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as X

from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData, BubbleChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from pptx_extract import (
    NS, ExtractionConfig, extract_pptx, iter_pptx_reports, write_outputs,
    enrich_descriptions, local_path, _chart_data,
)


def rewrite_zip(blob, changes):
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    parts.update(changes(parts))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return output.getvalue()


def make_demo():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Quarterly product revenue (USD millions)"
    series = CategoryChartData()
    series.categories = ["Q1", "Q2", "Q3"]
    series.add_series("Product A", [12, 15, 18])
    series.add_series("Product B", [10, 11, 13])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.3), Inches(8), Inches(4), series).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Revenue by quarter"
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = "USD millions"
    slide.notes_slide.notes_text_frame.text = "Illustrative data only. Calendar quarters."
    slide2 = deck.slides.add_slide(deck.slide_layouts[5])
    slide2.shapes.title.text = "Territory results"
    table = slide2.shapes.add_table(3, 3, Inches(1), Inches(1.5), Inches(8), Inches(2)).table
    for r, row in enumerate([["Region", "Revenue", "Growth"], ["East", "$12.5M", "8%"], ["West", "$9M", "6%"]]):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def change_workbook(blob):
    def patch(parts):
        part = next(p for p in parts if p.startswith("ppt/embeddings/"))
        wb = load_workbook(io.BytesIO(parts[part]))
        sheet = wb.create_sheet("Hidden assumptions")
        sheet.sheet_state = "hidden"
        sheet["A1"], sheet["B1"] = "Discount rate", "Formula result"
        sheet["A2"], sheet["B2"] = 0.125, "=2+3"
        sheet["A2"].number_format = "0.0%"
        sheet.row_dimensions[2].hidden = True
        sheet.column_dimensions["B"].hidden = True
        sheet.merge_cells("A4:B4")
        sheet["A4"] = "Merged note"
        out = io.BytesIO()
        wb.save(out)
        return {part: out.getvalue()}
    return rewrite_zip(blob, patch)


def make_multi_chart(titles=True, same_labels=False, grouped=False):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Regional performance"
    for i, label in enumerate(["East revenue", "West revenue"]):
        heading = slide.shapes.add_textbox(Inches(0.5 + i * 6), Inches(1.2), Inches(4.5), Inches(0.5))
        heading.text = label
        values = CategoryChartData()
        values.categories = ["Q1", "Q2"]
        values.add_series("Sales", [10 + i * 100, 20 + i * 100])
        frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5 + i * 6), Inches(2), Inches(4.5), Inches(3.5), values)
        if titles:
            frame.chart.has_title = True
            frame.chart.chart_title.text_frame.text = "Revenue" if same_labels else label
        if grouped:
            slide.shapes.add_group_shape([heading, frame])
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


class ExtractorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.demo = make_demo()

    def chart(self, report):
        return next(r for r in report["records"] if r["kind"] == "chart")

    def test_real_chart_table_notes_and_workbook(self):
        report = extract_pptx(self.demo)
        self.assertEqual(report["slide_count"], 2)
        chart = self.chart(report)
        self.assertEqual(chart["data"]["rows"][0]["value"], 12)
        self.assertEqual(chart["data"]["rows"][0]["category"], "Q1")
        self.assertIn("USD millions", chart["what_data_represents"])
        table = next(r for r in report["records"] if r["kind"] == "table")
        self.assertEqual(table["data"]["rows"][1], ["East", "$12.5M", "8%"])
        notes = next(r for r in report["records"] if r["kind"] == "speaker_notes")
        self.assertIn("Illustrative data", notes["data"]["text"])
        workbook = next(r for r in report["records"] if r["kind"] == "embedded_worksheet")
        self.assertEqual(workbook["slide_number"], 1)
        self.assertEqual(workbook["data"]["rows"][1][1], 12)
        self.assertFalse(any(r["kind"] == "extraction_issue" for r in report["records"]))
        json.dumps(report, allow_nan=False)

    def test_sparse_cache_does_not_shift_values(self):
        def patch(parts):
            root = X.fromstring(parts["ppt/charts/chart1.xml"])
            cache = root.find(".//c:ser/c:val/c:numRef/c:numCache", NS)
            cache.remove(cache.findall("c:pt", NS)[1])
            return {"ppt/charts/chart1.xml": X.tostring(root)}
        chart = self.chart(extract_pptx(rewrite_zip(self.demo, patch)))
        rows = chart["data"]["rows"]
        self.assertEqual([r["value"] for r in rows[:3]], [12, None, 18])
        self.assertEqual(rows[2]["category"], "Q3")
        self.assertTrue(chart["is_uncertain"])
        self.assertIn("MISSING_CHART_POINTS", [f["code"] for f in chart["uncertainty_reasons"]])

    def test_presentation_order_not_slide_filename(self):
        def patch(parts):
            root = X.fromstring(parts["ppt/presentation.xml"])
            slides = root.find("p:sldIdLst", NS)
            slides[:] = list(reversed(list(slides)))
            return {"ppt/presentation.xml": X.tostring(root)}
        report = extract_pptx(rewrite_zip(self.demo, patch))
        self.assertEqual(self.chart(report)["slide_number"], 2)
        self.assertEqual(next(r for r in report["records"] if r["kind"] == "table")["slide_number"], 1)

    def test_hidden_sheets_formula_and_format(self):
        report = extract_pptx(change_workbook(self.demo))
        sheet = next(r for r in report["records"] if r.get("sheet_name") == "Hidden assumptions")
        data = sheet["data"]
        self.assertEqual(data["sheet_state"], "hidden")
        self.assertEqual(data["rows"][1], [0.125, None])
        self.assertEqual(data["formulas"][0]["formula"], "=2+3")
        self.assertIn("lack cached results", " ".join(sheet["warnings"]))
        self.assertIn("MISSING_FORMULA_RESULT", [f["code"] for f in sheet["uncertainty_reasons"]])
        self.assertEqual(data["hidden_rows"], [2])
        self.assertEqual(data["merged_ranges"], ["A4:B4"])
        self.assertEqual(data["hidden_column_ranges"], [{"min": 2, "max": 2}])
        report = extract_pptx(change_workbook(self.demo), config=ExtractionConfig(include_hidden_sheets=False))
        self.assertFalse(any(r.get("sheet_name") == "Hidden assumptions" for r in report["records"]))

    def test_formula_cached_value_is_retained(self):
        blob = change_workbook(self.demo)
        def patch(parts):
            name = next(p for p in parts if p.startswith("ppt/embeddings/"))
            def patch_sheet(inner):
                node = X.fromstring(inner["xl/worksheets/sheet2.xml"])
                cell = next(c for c in node.findall("x:sheetData/x:row/x:c", NS) if c.get("r") == "B2")
                cell.find("x:v", NS).text = "5"
                return {"xl/worksheets/sheet2.xml": X.tostring(node)}
            return {name: rewrite_zip(parts[name], patch_sheet)}
        report = extract_pptx(rewrite_zip(blob, patch))
        sheet = next(r for r in report["records"] if r.get("sheet_name") == "Hidden assumptions")
        self.assertEqual(sheet["data"]["rows"][1][1], 5)
        self.assertEqual(sheet["data"]["formulas"][0]["cached_value"], 5)

    def test_external_links_are_reported_without_fetch(self):
        def patch(parts):
            name = "ppt/charts/_rels/chart1.xml.rels"
            root = X.fromstring(parts[name])
            X.SubElement(root, "Relationship", Id="rIdExternal", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink", Target="https://example.invalid/missing.xlsx", TargetMode="External")
            return {name: X.tostring(root)}
        report = extract_pptx(rewrite_zip(self.demo, patch))
        links = [r for r in report["records"] if r["kind"] == "external_link"]
        self.assertEqual(len(links), 1)
        self.assertEqual(links[0]["data"]["target"], "https://example.invalid/missing.xlsx")

    def test_missing_cache_has_warning_and_workbook(self):
        def patch(parts):
            root = X.fromstring(parts["ppt/charts/chart1.xml"])
            ref = root.find(".//c:ser/c:val/c:numRef", NS)
            ref.remove(ref.find("c:numCache", NS))
            return {"ppt/charts/chart1.xml": X.tostring(root)}
        report = extract_pptx(rewrite_zip(self.demo, patch))
        chart = self.chart(report)
        self.assertEqual(chart["status"], "partial")
        self.assertIsNone(chart["data"]["rows"][0]["value"])
        self.assertIn("no saved cache", " ".join(chart["warnings"]))
        self.assertTrue(any(r["kind"] == "embedded_worksheet" for r in report["records"]))

    def test_shared_workbook_provenance(self):
        def patch(parts):
            name = "ppt/slides/_rels/slide2.xml.rels"
            root = X.fromstring(parts[name])
            target = "../embeddings/" + next(p.rsplit("/", 1)[1] for p in parts if p.startswith("ppt/embeddings/"))
            X.SubElement(root, "Relationship", Id="rIdShared", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/package", Target=target)
            return {name: X.tostring(root)}
        report = extract_pptx(rewrite_zip(self.demo, patch))
        sheet = next(r for r in report["records"] if r["kind"] == "embedded_worksheet")
        self.assertIsNone(sheet["slide_number"])
        self.assertEqual({s["slide_number"] for s in sheet["data"]["source_references"]}, {1, 2})

    def test_scatter_and_bubble(self):
        for kind, cls in [(XL_CHART_TYPE.XY_SCATTER, XyChartData), (XL_CHART_TYPE.BUBBLE, BubbleChartData)]:
            with self.subTest(kind=kind):
                deck = Presentation()
                slide = deck.slides.add_slide(deck.slide_layouts[5])
                data = cls()
                series = data.add_series("Response")
                if cls is BubbleChartData:
                    series.add_data_point(1, 2, 3)
                    series.add_data_point(4, 5, 6)
                else:
                    series.add_data_point(1, 2)
                    series.add_data_point(4, 5)
                slide.shapes.add_chart(kind, Inches(1), Inches(1), Inches(7), Inches(4), data)
                out = io.BytesIO()
                deck.save(out)
                row = self.chart(extract_pptx(out.getvalue()))["data"]["rows"][1]
                self.assertEqual((row["x"], row["y"]), (4, 5))
                if cls is BubbleChartData:
                    self.assertEqual(row["bubble_size"], 6)

    def test_multilevel_categories_keep_sparse_levels(self):
        xml = f'''<c:chartSpace xmlns:c="{NS['c']}"><c:chart><c:plotArea><c:barChart><c:ser>
        <c:tx><c:v>Sales</c:v></c:tx><c:cat><c:multiLvlStrRef><c:f>Sheet1!A1:B2</c:f><c:multiLvlStrCache><c:ptCount val="2"/>
        <c:lvl><c:pt idx="0"><c:v>Q1</c:v></c:pt><c:pt idx="1"><c:v>Q2</c:v></c:pt></c:lvl>
        <c:lvl><c:pt idx="0"><c:v>2025</c:v></c:pt></c:lvl></c:multiLvlStrCache></c:multiLvlStrRef></c:cat>
        <c:val><c:numLit><c:ptCount val="2"/><c:pt idx="0"><c:v>1.20</c:v></c:pt><c:pt idx="1"><c:v>#N/A</c:v></c:pt></c:numLit></c:val>
        </c:ser></c:barChart></c:plotArea></c:chart></c:chartSpace>'''
        data, _ = _chart_data(X.fromstring(xml), ExtractionConfig())
        self.assertEqual(data["rows"][0]["category_levels"], ["Q1", "2025"])
        self.assertEqual(data["rows"][1]["category_levels"], ["Q2", None])
        self.assertEqual(data["rows"][0]["value_raw"], "1.20")
        self.assertEqual(data["rows"][1]["value"], "#N/A")

    def test_limits_emit_issue_not_silent_truncation(self):
        report = extract_pptx(self.demo, config=ExtractionConfig(max_sheet_cells=2))
        self.assertTrue(any("max_sheet_cells" in r["what_data_represents"] for r in report["records"] if r["kind"] == "extraction_issue"))
        with self.assertRaises(ValueError):
            extract_pptx(self.demo, config=ExtractionConfig(max_file_bytes=10))

    def test_hidden_slide_switch_and_orphan_embedding(self):
        def patch(parts):
            root = X.fromstring(parts["ppt/slides/slide1.xml"])
            root.set("show", "0")
            return {"ppt/slides/slide1.xml": X.tostring(root), "ppt/embeddings/legacy.bin": b"legacy-data"}
        blob = rewrite_zip(self.demo, patch)
        report = extract_pptx(blob)
        self.assertTrue(self.chart(report)["hidden_slide"])
        self.assertTrue(any(r["kind"] == "unsupported_embedding" for r in report["records"]))
        report = extract_pptx(blob, config=ExtractionConfig(include_hidden_slides=False, include_unreferenced_embeddings=False))
        self.assertFalse(any(r["kind"] == "chart" for r in report["records"]))

    def test_invalid_file_dtd_and_batch_isolation(self):
        with self.assertRaises(zipfile.BadZipFile):
            extract_pptx(b"not a zip")
        def patch(parts):
            return {"ppt/presentation.xml": b'<!DOCTYPE p [<!ENTITY bad "x">]><p/>'}
        with self.assertRaises(Exception) as caught:
            extract_pptx(rewrite_zip(self.demo, patch))
        self.assertIn("DTDForbidden", type(caught.exception).__name__)
        with tempfile.TemporaryDirectory() as tmp:
            good, bad = Path(tmp) / "good.pptx", Path(tmp) / "bad.pptx"
            good.write_bytes(self.demo)
            bad.write_bytes(b"bad")
            outputs = list(iter_pptx_reports([bad, good]))
            self.assertEqual([item["ok"] for item in outputs], [False, True])

    def test_paths_json_and_ids(self):
        self.assertEqual(local_path("dbfs:/Volumes/c/s/v/f.pptx").as_posix(), "/Volumes/c/s/v/f.pptx")
        self.assertEqual(local_path("dbfs:/old/f.pptx").as_posix(), "/dbfs/old/f.pptx")
        with self.assertRaises(ValueError):
            local_path("s3://bucket/file.pptx")
        report = extract_pptx(self.demo, source_name="demo.pptx")
        report2 = extract_pptx(self.demo, source_name="demo.pptx")
        self.assertEqual([r["record_id"] for r in report["records"]], [r["record_id"] for r in report2["records"]])
        with tempfile.TemporaryDirectory() as tmp:
            paths = write_outputs(report, tmp)
            reloaded = json.loads(Path(paths["report_json"]).read_text())
            self.assertEqual(reloaded, report)
            lines = Path(paths["records_jsonl"]).read_text().splitlines()
            self.assertEqual(len(lines), report["record_count"])

    def test_enrichment_does_not_mutate_data_or_original(self):
        report = extract_pptx(self.demo)
        original = copy.deepcopy(report)
        def describe(record):
            record["data"].clear()
            return "Illustrative quarterly product revenue, measured in USD millions."
        enriched = enrich_descriptions(report, describe)
        self.assertEqual(report, original)
        self.assertEqual(self.chart(enriched)["data"], self.chart(original)["data"])
        self.assertEqual(self.chart(enriched)["description_method"], "user_callback_inferred")
        self.assertTrue(self.chart(enriched)["needs_review"])
        self.assertIn("INFERRED_DESCRIPTION", [f["code"] for f in self.chart(enriched)["uncertainty_reasons"]])
        self.assertEqual(enriched["uncertain_record_count"], sum(r["is_uncertain"] for r in enriched["records"]))

    def test_side_by_side_charts_have_separate_values_labels_and_positions(self):
        report = extract_pptx(make_multi_chart())
        charts = [r for r in report["records"] if r["kind"] == "chart"]
        self.assertEqual(len(charts), 2)
        self.assertEqual([c["chart_index_on_slide"] for c in charts], [1, 2])
        self.assertEqual([c["chart_count_on_slide"] for c in charts], [2, 2])
        self.assertEqual([c["position_on_slide"] for c in charts], ["middle-left", "middle-right"])
        self.assertNotEqual(charts[0]["record_id"], charts[1]["record_id"])
        self.assertNotEqual(charts[0]["shape_id"], charts[1]["shape_id"])
        self.assertNotEqual(charts[0]["source_part"], charts[1]["source_part"])
        self.assertEqual([r["value"] for r in charts[0]["data"]["rows"]], [10, 20])
        self.assertEqual([r["value"] for r in charts[1]["data"]["rows"]], [110, 120])
        self.assertIn("East revenue", charts[0]["what_data_represents"])
        self.assertNotIn("West revenue", charts[0]["what_data_represents"])
        self.assertEqual([c["needs_review"] for c in charts], [False, False])
        for c in charts:
            self.assertEqual(c["quality"]["source_freshness"], "unverified_saved_values")
        worksheets = [r for r in report["records"] if r["kind"] == "embedded_worksheet"]
        self.assertEqual([r["data"]["source_references"][0]["chart_index_on_slide"] for r in worksheets], [1, 2])

    def test_nearby_headings_are_separate_and_explicitly_uncertain(self):
        charts = [r for r in extract_pptx(make_multi_chart(titles=False))["records"] if r["kind"] == "chart"]
        for chart, label in zip(charts, ["East revenue", "West revenue"]):
            self.assertTrue(chart["is_uncertain"])
            self.assertEqual(chart["quality"]["data_extraction"], "no_detected_issue")
            self.assertEqual(chart["quality"]["meaning"], "ambiguous")
            flags = [f["code"] for f in chart["uncertainty_reasons"]]
            self.assertIn("NEARBY_LABEL_UNVERIFIED", flags)
            self.assertIn("MULTIPLE_CHARTS_WITHOUT_LOCAL_LABEL", flags)
            self.assertEqual([n["text"] for n in chart["description_evidence"]["nearby_text_candidates"]], [label])

    def test_identical_chart_labels_are_flagged_without_merging_data(self):
        charts = [r for r in extract_pptx(make_multi_chart(same_labels=True))["records"] if r["kind"] == "chart"]
        self.assertEqual(len(charts), 2)
        self.assertTrue(all("DUPLICATE_CHART_LABELS" in [f["code"] for f in r["uncertainty_reasons"]] for r in charts))
        self.assertNotEqual(charts[0]["data"]["rows"], charts[1]["data"]["rows"])

    def test_grouped_chart_positions_and_rotated_group_fallback(self):
        blob = make_multi_chart(grouped=True)
        charts = [r for r in extract_pptx(blob)["records"] if r["kind"] == "chart"]
        self.assertEqual([c["position_on_slide"] for c in charts], ["middle-left", "middle-right"])
        def patch(parts):
            root = X.fromstring(parts["ppt/slides/slide1.xml"])
            xfrm = root.find(".//p:grpSp/p:grpSpPr/a:xfrm", NS)
            xfrm.set("rot", "5400000")
            return {"ppt/slides/slide1.xml": X.tostring(root)}
        charts = [r for r in extract_pptx(rewrite_zip(blob, patch))["records"] if r["kind"] == "chart"]
        self.assertIsNone(charts[0]["bounding_box"])
        self.assertEqual(charts[0]["data"]["rows"][0]["value"], 10)
        self.assertIn("CHART_POSITION_UNRESOLVED", [f["code"] for f in charts[0]["uncertainty_reasons"]])
        self.assertEqual(charts[1]["position_on_slide"], "middle-right")

    def test_mismatched_counts_and_error_values_have_uncertainty_reasons(self):
        def patch(parts):
            root = X.fromstring(parts["ppt/charts/chart1.xml"])
            cache = root.find(".//c:ser/c:val/c:numRef/c:numCache", NS)
            cache.find("c:ptCount", NS).set("val", "4")
            cache.find("c:pt/c:v", NS).text = "#N/A"
            return {"ppt/charts/chart1.xml": X.tostring(root)}
        chart = self.chart(extract_pptx(rewrite_zip(self.demo, patch)))
        flags = [f["code"] for f in chart["uncertainty_reasons"]]
        self.assertIn("MISMATCHED_POINT_COUNTS", flags)
        self.assertIn("NONNUMERIC_CHART_VALUE", flags)
        self.assertEqual(chart["data"]["rows"][0]["value"], "#N/A")

    def test_malformed_geometry_does_not_drop_chart_data(self):
        def patch(parts):
            root = X.fromstring(parts["ppt/slides/slide1.xml"])
            root.find(".//p:graphicFrame/p:xfrm/a:off", NS).set("x", "invalid")
            return {"ppt/slides/slide1.xml": X.tostring(root)}
        report = extract_pptx(rewrite_zip(make_multi_chart(), patch))
        charts = [r for r in report["records"] if r["kind"] == "chart"]
        self.assertEqual(len(charts), 2)
        self.assertEqual(charts[0]["data"]["rows"][0]["value"], 10)
        self.assertTrue(all(c["needs_review"] for c in charts))


if __name__ == "__main__":
    unittest.main()
